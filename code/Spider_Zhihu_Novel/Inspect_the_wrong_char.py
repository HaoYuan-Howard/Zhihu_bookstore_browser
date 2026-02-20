from docx import Document

def scan_special_chars(docx_path):
    # 加载Word文档
    doc = Document(docx_path)

    suspicious_chars = {}  # 统计用 {字符: 出现次数}

    def check_text(text):
        for ch in text:
            codepoint = ord(ch)
            # 检查康熙部首和中日韩部首补充区
            if (0x2F00 <= codepoint <= 0x2FD5) or (0x2E80 <= codepoint <= 0x2EFF):
                suspicious_chars[ch] = suspicious_chars.get(ch, 0) + 1

    # 遍历段落
    for para in doc.paragraphs:
        check_text(para.text)

    # 遍历表格
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                check_text(cell.text)

    return suspicious_chars


if __name__ == "__main__":
    input_file = r"E:\practice_group\generate_novel\material\novel\后续\发现自己是恶毒女配后.docx"
    result = scan_special_chars(input_file)

    if result:
        print("⚠️ 发现可疑字符：")
        for ch, count in result.items():
            print(f"字符: {ch} (U+{ord(ch):04X}) -> 出现 {count} 次")
    else:
        print("✅ 文档中没有发现康熙部首或部首补充区字符。")
